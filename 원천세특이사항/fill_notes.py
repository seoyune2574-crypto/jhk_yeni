"""특이사항 엑셀(원천세) 내용을 원천세신고리스트 '인수인계파일 원천세 특이사항' 칸에
사업자등록번호 기준으로 채워넣음 (2026-09-29)
- 신고리스트에 있는 사업장만 채움
- 특이사항 파일의 병합셀은 병합범위 모든 행에 같은 내용으로 적용
사용: python fill_notes.py 원천세신고리스트.xlsx 특이사항.xlsx 결과.xlsx
"""
import re, sys
import openpyxl
from openpyxl.styles import Alignment

n = lambda s: re.sub(r'\D', '', str(s)) if s is not None else ''
target_path, notes_path, out = sys.argv[1:4]

# 특이사항 파일: C=사업자등록번호, E=특이사항(원천세)
ns = openpyxl.load_workbook(notes_path).active
merged = {}
for rng in ns.merged_cells.ranges:
    if rng.min_col <= 5 <= rng.max_col:
        v = ns.cell(rng.min_row, 5).value
        for r in range(rng.min_row, rng.max_row + 1):
            merged[r] = v
notes = {}
for r in range(3, ns.max_row + 1):
    k = n(ns.cell(r, 3).value)
    if len(k) != 10:
        continue
    v = merged.get(r, ns.cell(r, 5).value)
    v = str(v).strip() if v is not None else ''
    if v and v not in notes.get(k, []):
        notes.setdefault(k, []).append(v)

# 원천세신고리스트: D=사업자번호, J=인수인계파일 원천세 특이사항
wb = openpyxl.load_workbook(target_path)
ws = wb.active
filled, empty, missing = [], [], []
for r in range(2, ws.max_row + 1):
    k = n(ws.cell(r, 4).value)
    if not k:
        continue
    name = ws.cell(r, 3).value
    if k not in notes and k not in {n(ns.cell(i, 3).value) for i in range(3, ns.max_row + 1)}:
        missing.append((r, name, ws.cell(r, 4).value))
    elif k in notes:
        c = ws.cell(r, 10, '\n'.join(notes[k]))
        c.alignment = Alignment(wrap_text=True, vertical='center')
        filled.append((r, name))
    else:
        empty.append((r, name))
wb.save(out)
print(f'채움 {len(filled)} / 특이사항 빈칸 {len(empty)} / 특이사항파일에 없음 {len(missing)}')
print('빈칸:', empty)
print('없음:', missing)
