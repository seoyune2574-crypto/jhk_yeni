"""거래처 관리 인폼의 사업자번호와 매칭되는 예정고지세액 현황 행에
담당자(F열) '정서연' + 기존 작업과 같은 음영(#FCE5CD)을 표시한다.

사용: python mark.py 인폼.xlsx 대상.xlsx 결과.xlsx
"""
import re
import sys

import openpyxl
from openpyxl.styles import PatternFill

NAME = '정서연'
FILL = PatternFill(fill_type='solid', fgColor='FFFCE5CD')
# 시트명: (데이터 시작행, 사업자번호 열 index) — 폐업/해지 시트는 제외
SOURCES = {'기장거래처': (5, 10), '신고대리': (3, 11)}


def brn(v):
    if v is None:
        return ''
    if isinstance(v, float):
        return str(int(v)).zfill(10)
    return re.sub(r'\D', '', str(v))


def main(inform_path, target_path, out_path):
    wb = openpyxl.load_workbook(inform_path, read_only=True)
    mine = set()
    for name, (start, col) in SOURCES.items():
        for row in wb[name].iter_rows(min_row=start, values_only=True):
            b = brn(row[col])
            if len(b) == 10:
                mine.add(b)

    wt = openpyxl.load_workbook(target_path)
    ws = wt.active
    added = already = 0
    for row in ws.iter_rows(min_row=2):
        if brn(row[4].value) not in mine:
            continue
        cell = row[5]
        if cell.value == NAME:
            already += 1
        else:
            added += 1
        cell.value = NAME
        cell.fill = FILL
    wt.save(out_path)
    print(f'인폼 사업자번호 {len(mine)}개 / 매칭 {added + already}건 (신규 {added}, 기존 {already})')


if __name__ == '__main__':
    main(*sys.argv[1:4])

# 결과물 보관: Google Drive '클로드' 폴더
# https://drive.google.com/drive/folders/1qWzDPx_IfHt5zsWCEWE1jRyEqtsXonfe
