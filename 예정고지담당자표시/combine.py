"""거래처 관리 인폼(기장거래처·신고대리) 전체 거래처의 상호·사업자번호·대표자·주민번호 우측에
26년 2기 예정고지세액 현황 행을 사업자번호로 붙인 CSV를 만든다.
폐업·해지 시트는 다른 담당자 거래처라 제외. 주민번호가 들어가므로 결과 파일은 저장소에 커밋하지 않는다.

사용: python combine.py 인폼.xlsx 예정고지.xlsx 결과.csv
결과물 보관: Google Drive '클로드' 폴더
"""
import collections
import csv
import re
import sys

import openpyxl

# 시트명: (시작행, 대표자, 사업자번호, 주민번호, 상호 열 index)
SOURCES = {'기장거래처': (5, 9, 10, 11, 12), '신고대리': (3, 10, 11, 12, 13)}


def brn(v):
    if v is None:
        return ''
    return re.sub(r'\D', '', str(int(v)) if isinstance(v, float) else str(v))


def text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def main(inform_path, target_path, out_path):
    ws = openpyxl.load_workbook(target_path).active
    hdr = [c.value.replace('\n', ' ') for c in ws[1]]
    bcol = next(i for i, h in enumerate(hdr) if h.startswith('사업자 등록번호'))
    tgt = collections.defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        tgt[brn(r[bcol])].append(r)

    wb = openpyxl.load_workbook(inform_path, read_only=True)
    with open(out_path, 'w', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(['구분', '상호', '사업자번호', '대표자', '주민등록번호'] + hdr)
        for name, (start, rep, bc, jm, nm) in SOURCES.items():
            for r in wb[name].iter_rows(min_row=start, values_only=True):
                r = list(r) + [None] * 30
                if all(r[i] in (None, '') for i in (rep, bc, jm, nm)):
                    continue
                b = brn(r[bc])
                hit = tgt.get(b) if len(b) == 10 else None
                extra = [text(v) for v in hit[0]] if hit else [''] * len(hdr)
                w.writerow([name, text(r[nm]).strip(), text(r[bc]), text(r[rep]), text(r[jm])] + extra)


if __name__ == '__main__':
    main(*sys.argv[1:4])
